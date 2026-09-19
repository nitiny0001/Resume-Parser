from io import BytesIO

from docx import Document

from app.services.text_extractor import DOCX_TYPE, extract_text


def test_extract_docx_text():
    document = Document()
    document.add_paragraph("Nitin Kumar")
    document.add_paragraph("Python FastAPI PostgreSQL")
    stream = BytesIO()
    document.save(stream)

    text, pages, method = extract_text(stream.getvalue(), DOCX_TYPE)

    assert "Nitin Kumar" in text
    assert "FastAPI" in text
    assert pages == 1
    assert method == "python-docx"
